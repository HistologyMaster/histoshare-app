#!/usr/bin/env python3
"""export.py <final_dir> : renders figures and exports the Markdown review to DOCX, HTML and PDF (pandoc + LibreOffice + headless Chromium)."""
import os, subprocess, sys, shutil, re
D = sys.argv[1]
WS = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
fig = os.path.join(D, 'figuras'); os.makedirs(fig, exist_ok=True)
svg = os.path.join(fig, 'fig1_prisma.svg'); png = os.path.join(fig, 'fig1_prisma.png')
subprocess.run([sys.executable, '-I', os.path.join(WS, 'tools', 'fig_prisma.py'), svg], check=True)
CH = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'
tmp = os.path.join(fig, '_t.png')
subprocess.run([CH, '--headless=new', '--no-sandbox', '--disable-gpu', '--hide-scrollbars', '--window-size=1000,1200', f'--screenshot={tmp}', 'file://' + os.path.abspath(svg)], check=True, capture_output=True, timeout=120)
subprocess.run(['convert', tmp, '-crop', '1000x980+0+0', '+repage', png], check=True); os.remove(tmp)
md = os.path.join(D, 'revision_paraphysis_cerebri.md')
base = os.path.join(D, 'revision_paraphysis_cerebri')
meta = ['-M', 'lang=es', '-M', 'pagetitle=Paráfisis cerebral (paraphysis cerebri): revisión sistemática narrativa']
# DOCX
subprocess.run(['pandoc', md, '-f', 'markdown+pipe_tables+smart', '-o', base + '.docx', '--resource-path', D] + meta, check=True)
# post-process tables with python-docx
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document(base + '.docx')
def borders(t):
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + k); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '9AA5B1'); b.append(e)
    tblPr.append(b)
def shade(cell):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'EEF3F8'); tcPr.append(sh)
for t in doc.tables:
    borders(t)
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            if ri == 0: shade(c)
            for p_ in c.paragraphs:
                for r in p_.runs:
                    r.font.size = Pt(8)
                    if ri == 0: r.font.bold = True
doc.save(base + '.docx')
# HTML (self-contained)
css = os.path.join(WS, 'tools', 'review.css')
open(css, 'w').write('body{max-width:920px;margin:2rem auto;padding:0 1rem;font:16px/1.55 Georgia,serif;color:#1f2933}table{border-collapse:collapse;font-size:.82rem;margin:1rem 0;width:100%}th,td{border:1px solid #c9d2dc;padding:.3rem .45rem;vertical-align:top}th{background:#eef3f8}blockquote{border-left:4px solid #a9742c;background:#fbf3e6;margin:1rem 0;padding:.6rem 1rem}h1{font-size:1.6rem}h2{margin-top:2rem;border-bottom:1px solid #c9d2dc}img{max-width:100%}code{font-size:.85em}')
subprocess.run(['pandoc', md, '-f', 'markdown+pipe_tables+smart', '-s', '--toc', '--toc-depth=2', '--embed-resources', '--standalone', '-c', css, '--resource-path', D, '-o', base + '.html'] + meta, check=True)
# PDF via LibreOffice from the DOCX
subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', D, base + '.docx'], check=True, capture_output=True, timeout=300)
for f in os.listdir(D): print(f, os.path.getsize(os.path.join(D, f)) if os.path.isfile(os.path.join(D, f)) else 'dir')
